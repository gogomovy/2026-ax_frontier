"""시연 자료 엑셀 작성기 (Phase I · 설계 §25.2) — 실무에서 내려받는 모양 그대로.

- 실무형 표: 제목(병합 · 굵게) · 기준/단위/작성 줄 · 빈 줄 · 머리글(채움 · 테두리 · 틀 고정) · 천 단위 · 소수 서식 ·
  코드는 문자열 셀(앞자리 0 보존) · 합계 줄 · 주석(`※ 시연용 가상 자료`)
- 원본 보고서 양식: 2633 손익(계층 항목 · 들여쓰기 · `N월 누계`), 2020 월별 여신(머리글 블록 · 과목구분 · 1~N월 · `[1,000,000] 원`),
  7790 충당금(다단 병합 머리글 · X 열 총계 · 합계 행)
- 같은 입력이면 같은 바이트(문서 속성 · 압축 항목 시각 고정) — 저장소 사본과 생성 결과를 비교할 수 있게
"""
from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FIXED = datetime(2026, 9, 1, 9, 0, 0)
NOTICE = "※ 시연용 가상 자료 — 실제 은행 수치가 아닙니다(NH 경영 의사결정 플랫폼 업로드 시연용으로 생성)."
HEAD_FILL = PatternFill("solid", fgColor="E3EFE6")
TOTAL_FILL = PatternFill("solid", fgColor="F3F4F1")
THIN = Side(style="thin", color="9AA79E")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE = Font(bold=True, size=14)
META = Font(size=9, color="555555")
BOLD = Font(bold=True)


def finish(wb: Workbook) -> bytes:
    """저장 + 결정적 바이트(작성 시각 · 압축 항목 시각 고정)."""
    wb.properties.creator = "NH 경영 의사결정 플랫폼 · 시연 자료 생성기"
    wb.properties.lastModifiedBy = wb.properties.creator
    wb.properties.created = wb.properties.modified = FIXED
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":  # openpyxl 이 저장 시각으로 덮어쓰는 수정 시각을 고정
                import re

                data = re.sub(rb"<dcterms:modified([^>]*)>[^<]*</dcterms:modified>", b"<dcterms:modified\\1>2026-09-01T09:00:00Z</dcterms:modified>", data)
            z.writestr(zi, data)
    return out.getvalue()


@dataclass
class Col:
    header: str  # 머리글(데이터셋 정의의 라벨 · 별칭과 맞춘다)
    key: str  # 행 dict 키
    fmt: str = "text"  # text · int · num0 · num1 · num2 · num3 · num4 · pct2 · date · bool
    width: float | None = None


_FMT = {"int": "#,##0", "num0": "#,##0", "num1": "#,##0.0", "num2": "#,##0.00", "num3": "#,##0.000", "num4": "#,##0.0000", "pct2": "0.00", "text": "@"}


def _value(v: Any, fmt: str) -> Any:
    if v is None:
        return None
    if fmt == "text":
        return str(v)
    if fmt == "bool":
        return "Y" if v else "N"
    if fmt == "date":
        return str(v)
    return v


def _width(header: str, values: Sequence[Any]) -> float:
    def w(s: Any) -> int:
        t = "" if s is None else (f"{s:,}" if isinstance(s, (int, float)) and not isinstance(s, bool) else str(s))
        return sum(2 if ord(ch) > 0x2E80 else 1 for ch in t)

    return min(48, max(8, w(header) + 2, *(w(v) + 2 for v in values[:200])))


def table(
    title: str,
    meta: list[str],
    cols: list[Col],
    rows: list[dict[str, Any]],
    *,
    sheet: str = "Sheet1",
    total: dict[str, Any] | None = None,
    notes: list[str] | None = None,
) -> bytes:
    """실무형 표 한 장."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    n = len(cols)
    ws.cell(row=1, column=1, value=title).font = TITLE
    if n > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(n, 8))
    r = 2
    for line in meta:
        ws.cell(row=r, column=1, value=line).font = META
        r += 1
    r += 1  # 빈 줄
    hr = r
    for c, col in enumerate(cols, start=1):
        cell = ws.cell(row=hr, column=c, value=col.header)
        cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        r += 1
        for c, col in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=c, value=_value(row.get(col.key), col.fmt))
            cell.border = BOX
            if col.fmt in _FMT:
                cell.number_format = _FMT[col.fmt]
            if col.fmt not in ("text", "bool", "date"):
                cell.alignment = Alignment(horizontal="right")
    if total:
        r += 1
        for c, col in enumerate(cols, start=1):
            v = total.get(col.key)
            cell = ws.cell(row=r, column=c, value=_value(v, col.fmt) if v is not None else None)
            cell.font, cell.fill, cell.border = BOLD, TOTAL_FILL, BOX
            if col.fmt in _FMT and col.fmt != "text":
                cell.number_format = _FMT[col.fmt]
    r += 2
    for line in [*(notes or []), NOTICE]:
        ws.cell(row=r, column=1, value=line).font = META
        r += 1
    for c, col in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(c)].width = col.width or _width(col.header, [row.get(col.key) for row in rows])
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    return finish(wb)


def csv_bytes(cols: list[Col], rows: list[dict[str, Any]]) -> bytes:
    """CSV(UTF-8 BOM, 엑셀에서 바로 열리게) — 머리글 1줄 + 행. 시스템 추출 파일 모양."""
    import csv

    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow([c.header for c in cols])
    for row in rows:
        out = []
        for c in cols:
            v = row.get(c.key)
            if v is None:
                out.append("")
            elif c.fmt == "bool":
                out.append("1" if v else "0")
            elif isinstance(v, float) and v.is_integer() and c.fmt in ("int", "num0"):
                out.append(str(int(v)))
            else:
                out.append(str(v))
        w.writerow(out)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


# ═════════════════════════ 원본 보고서 양식 ═════════════════════════
def _indent(ws, r: int, c: int, value: Any, level: int) -> None:
    cell = ws.cell(row=r, column=c, value=value)
    if level:
        cell.alignment = Alignment(indent=level)


def report_2633(ym: str, items: list[tuple[str, int, float | None, float | None]], *, office: str = "00510(종합기획부)") -> bytes:
    """2633 상품부문별 손익계산서 — items: (항목명, 들여쓰기, 당월 누계, 전년 동기 누계). 단위 백만원."""
    y, m = ym[:4], int(ym[4:])
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "2633 상품부문별 손익계산서"
    ws["A1"].font = TITLE
    ny, nm = (int(y) + 1, 1) if m == 12 else (int(y), m + 1)
    ws["A2"] = f"조회일자 : {ny}-{nm:02d}-03 09:12:44"
    ws["A3"] = f"조회사무소 : {office}"
    ws["A4"] = f"기 준 일 : [{y}-{m:02d}-{_last_day(int(y), m):02d}], 사무소 : [전행], 단위 : 백만원"
    for a in ("A2", "A3", "A4"):
        ws[a].font = META
    hr = 6
    heads = ["항목명", "당월", "전월 누계", f"{m}월 누계", "전년동기", "증감률(%)"]
    for c, h in enumerate(heads, start=1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center")
    for i, (name, level, cum, prev) in enumerate(items):
        r = hr + 1 + i
        _indent(ws, r, 1, name, level)
        if cum is not None:
            month = round(cum / m)
            ws.cell(row=r, column=2, value=month).number_format = "#,##0"
            ws.cell(row=r, column=3, value=round(cum - month)).number_format = "#,##0"
            ws.cell(row=r, column=4, value=cum).number_format = "#,##0"
        if prev is not None:
            ws.cell(row=r, column=5, value=prev).number_format = "#,##0"
            if cum is not None and prev:
                ws.cell(row=r, column=6, value=round((cum - prev) / prev * 100, 2)).number_format = "0.00"
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = BOX
    r = hr + len(items) + 2
    ws.cell(row=r, column=1, value=NOTICE).font = META
    ws.column_dimensions["A"].width = 34
    for col in "BCDEF":
        ws.column_dimensions[col].width = 14
    ws.freeze_panes = ws.cell(row=hr + 1, column=2)
    return finish(wb)


def report_2020(year: int, months: int, tree: list[tuple[str, int, list[float | None]]], *, base_date: str, office: str = "00561(신용리스크관리부)") -> bytes:
    """2020 과목별 월별 여신현황 — tree: (과목명, 들여쓰기, 1~N월 잔액(백만원)). 단위 [1,000,000] 원."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "2020 과목별 월별 여신현황"
    ws["A1"].font = TITLE
    ws["A2"] = f"조회일자 : {base_date} 18:18:23"
    ws["A3"] = f"조회사무소 : {office}"
    ws["A5"] = "조회구분 : [전국], 사무소 : [], 과목 : [여신대분류], 유형 : []"
    ws["A6"] = f"기 준 일 : [{base_date}], 사무소명 : [], 구분 : [계리잔액], 조건 : [], 단위:명,좌, [1,000,000] 원"
    for a in ("A2", "A3", "A5", "A6"):
        ws[a].font = META
    hr = 8
    head = ws.cell(row=hr, column=1, value="과목구분")
    head.font, head.fill, head.border = BOLD, HEAD_FILL, BOX
    for mo in range(1, months + 1):
        cell = ws.cell(row=hr, column=1 + mo, value=f"{mo}월")
        cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center")
    for i, (name, level, vals) in enumerate(tree):
        r = hr + 1 + i
        _indent(ws, r, 1, name, level)
        ws.cell(row=r, column=1).border = BOX
        for mo, v in enumerate(vals[:months], start=1):
            cell = ws.cell(row=r, column=1 + mo, value=v)
            cell.number_format, cell.border = "#,##0", BOX
    r = hr + len(tree) + 2
    ws.cell(row=r, column=1, value=NOTICE).font = META
    ws.column_dimensions["A"].width = 30
    for mo in range(1, months + 1):
        ws.column_dimensions[get_column_letter(1 + mo)].width = 13
    ws.freeze_panes = ws.cell(row=hr + 1, column=2)
    return finish(wb)


def report_7790(ym: str, products: list[dict[str, Any]], total: float) -> bytes:
    """7790 상품별 충당금 산출현황 — 다단 머리글 · X 열 '손실충당금 및 충당부채 총계' · 합계 행. 단위 백만원.

    products: {code, name, borrowers, provision}
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "7790"
    ws["A1"] = "7790 상품별 충당금 산출현황 요약"
    ws["A1"].font = TITLE
    ws["A2"] = f"기준년월 : [{ym[:4]}.{ym[4:]}]"
    ws["A3"] = "단위 : 백만원 · 조회사무소 : 00561(신용리스크관리부)"
    for a in ("A2", "A3"):
        ws[a].font = META
    head = {"A4": "상품코드", "B4": "상품명", "C4": "차주수", "D4": "손실충당금", "G4": "충당부채", "H4": "Stage 구분", "X4": "손실충당금 및 충당부채 총계",
            "D5": "대출채권", "E5": "미수이자", "F5": "가지급금", "G5": "미사용약정", "H5": "Stage1", "I5": "Stage2", "J5": "Stage3"}
    ws.merge_cells("D4:F4")
    ws.merge_cells("H4:J4")
    for ref, v in head.items():
        cell = ws[ref]
        cell.value = v
        cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ref in ("A5", "B5", "C5", "X5"):
        ws[ref].fill, ws[ref].border = HEAD_FILL, BOX
    r = 6
    for p in products:
        v = float(p["provision"])
        loan, accr, adv = round(v * 0.91, 1), round(v * 0.05, 1), round(v * 0.01, 1)
        unused = round(v - loan - accr - adv, 1)
        s1, s2 = round(v * 0.38, 1), round(v * 0.41, 1)
        cells = {1: str(p["code"]), 2: p["name"], 3: p.get("borrowers"), 4: loan, 5: accr, 6: adv, 7: unused, 8: s1, 9: s2, 10: round(v - s1 - s2, 1), 24: p["provision"]}
        for c, val in cells.items():
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = BOX
            if c == 1:
                cell.number_format = "@"
            elif c >= 3:
                cell.number_format = "#,##0" if c == 3 else "#,##0.0"
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="합계").font = BOLD
    t = ws.cell(row=r, column=24, value=total)
    t.font, t.number_format = BOLD, "#,##0"
    ws.cell(row=r + 2, column=1, value=NOTICE).font = META
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["X"].width = 18
    ws.freeze_panes = "C6"
    return finish(wb)


def _last_day(y: int, m: int) -> int:
    import calendar

    return calendar.monthrange(y, m)[1]
