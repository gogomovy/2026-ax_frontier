"""업무 구성 파일 형식 (Phase O · BIZ-02 · Q215) — 엑셀(시트별 표) 또는 같은 구조의 JSON.

시드 대신 관리자가 올리는 업무 구조: 자료(원천 이름 · 필요한 열 · 별칭) · 지표 · 업무 흐름(단계 · 조건) · 화면 연결 · 화면 틀 · 규칙 ·
조직 · 화면 · 권한 · 설정 · 서술 에이전트 · 지식 컬렉션 · 온톨로지 · 정의서. 구조가 있는 칸(지표 입력 · 규칙 본문 · 화면 틀)은 JSON 글자.
"""
from __future__ import annotations

import io
import json
from typing import Any

FORMAT = "nhdip.business-config/v1"
CHUNK = 30_000  # 엑셀 칸 글자 수 한도(32,767) 아래 — 화면 틀은 여러 줄로 나눈다
MAX_BYTES = 20 * 1024 * 1024


class BundleError(ValueError):
    pass


# (시트, 묶음 키, [(머리글, 필드, 종류)]) — 종류: s 글자 · i 정수 · f 숫자 · b 참거짓 · j JSON · l 목록(쉼표 또는 JSON)
SHEETS: list[tuple[str, str, list[tuple[str, str, str]]]] = [
    ("자료", "sources", [("자료 이름", "name", "s"), ("필요한 열", "columns", "l"), ("이름 단서", "keywords", "l"), ("코드용 열 별칭(JSON)", "aliases", "j"),
                       ("온톨로지 연결(JSON)", "mapping", "j"), ("설명", "description", "s")]),
    ("조직", "orgs", [("코드", "code", "s"), ("이름", "name", "s"), ("유형", "unit_type", "s"), ("상위 코드", "parent_code", "s"), ("순서", "sort", "i")]),
    ("화면", "screens", [("코드", "code", "s"), ("주소", "path", "s"), ("이름", "label", "s"), ("메뉴 묶음", "section", "s"), ("종류", "kind", "s"), ("단계", "step", "i"),
                        ("건수 키", "count_key", "s"), ("활성 판정", "match", "s"), ("순서", "sort", "i"), ("관리자 전용", "admin_only", "b"), ("사용", "active", "b"),
                        ("화면 자원", "resources", "l")]),
    ("화면권한", "grants", [("조직", "org", "s"), ("화면", "screen", "s")]),
    ("설정", "settings", [("키", "key", "s"), ("값(JSON)", "value", "j"), ("반영", "mode", "s"), ("설명", "description", "s")]),
    ("규칙", "rules", [("코드", "code", "s"), ("이름", "name", "s"), ("종류", "kind", "s"), ("설명", "description", "s"), ("본문(JSON)", "body", "j"), ("공통 규칙에 더함", "merge", "b")]),
    ("에이전트", "agents", [("코드", "code", "s"), ("이름", "name", "s"), ("역할", "role", "s"), ("종류", "kind", "s"), ("그룹", "group_code", "s"), ("프로필", "profile", "s"),
                         ("프롬프트", "system", "s"), ("입력 틀", "user_template", "s")]),
    ("업무흐름", "workflows", [("코드", "code", "s"), ("이름", "name", "s"), ("설명", "description", "s"), ("산출물", "outputs", "l"), ("순서", "sort", "i"), ("사용", "enabled", "b")]),
    ("흐름단계", "_nodes", [("업무흐름", "workflow", "s"), ("순서", "seq", "i"), ("키", "key", "s"), ("이름", "name", "s"), ("구현", "impl", "s"), ("에이전트", "agent_code", "s"),
                         ("설정(JSON)", "config", "j"), ("사용", "enabled", "b")]),
    ("흐름조건", "_triggers", [("업무흐름", "workflow", "s"), ("종류", "kind", "s"), ("대상", "target", "s"), ("지연(초)", "debounce_sec", "i"), ("사용", "enabled", "b")]),
    ("지표", "metrics", [("코드", "code", "s"), ("이름", "label", "s"), ("산식", "formula", "s"), ("입력(JSON)", "args", "j"), ("단위", "unit", "s"), ("소수", "decimals", "i"),
                        ("반올림", "rounding", "s"), ("배율", "scale", "f"), ("업무흐름", "workflow_code", "s"), ("화면 칸", "figure_ids", "l"), ("설명", "description", "s"),
                        ("사용", "enabled", "b"), ("순서", "sort", "i")]),
    ("화면연결", "bindings", [("업무흐름", "workflow_code", "s"), ("화면 자료", "artifact_key", "s"), ("위치", "path", "s"), ("만드는 방법", "builder", "s"), ("설정(JSON)", "params", "j"),
                          ("사용", "enabled", "b")]),
    ("화면틀", "_artifacts", [("화면 자료", "key", "s"), ("조각", "part", "i"), ("내용(JSON 조각)", "chunk", "s")]),
    ("지식컬렉션", "collections", [("코드", "code", "s"), ("이름", "label", "s"), ("설명", "description", "s"), ("종류", "kind", "s"), ("담당 조직", "owner_orgs", "l"),
                              ("공유 조직", "shared_orgs", "l"), ("필요 화면 자원", "required_resources", "l"), ("공개 범위", "audience", "s"), ("포함 규칙(JSON)", "sources", "j"),
                              ("가중", "boost", "f"), ("사용", "enabled", "b"), ("순서", "sort", "i")]),
    ("온톨로지", "_ontology", [("종류", "kind", "s"), ("이름", "name", "s"), ("내용(JSON)", "body", "j")]),
    ("정의서", "contracts", [("표 코드", "table_code", "s"), ("표 이름", "table_name", "s"), ("부서", "dept", "s"), ("순서", "seq", "i"), ("컬럼", "column_name", "s"),
                          ("한글 이름", "kor_name", "s"), ("기본키", "is_pk", "b"), ("타입", "data_type", "s"), ("원천 설명", "source_desc", "s"), ("원천 컬럼", "source_column", "s"),
                          ("목업", "is_mock", "b"), ("기대 화면값", "expected_text", "s"), ("연결(JSON)", "binding", "j"), ("사용", "active", "b"), ("정의서 파일", "workbook", "s")]),
]


def _cell(v: Any, kind: str) -> Any:
    if v is None:
        return None
    if kind == "j":
        return json.dumps(v, ensure_ascii=False)
    if kind == "l":
        return ", ".join(str(x) for x in v) if all(isinstance(x, str) and "," not in x for x in v) else json.dumps(v, ensure_ascii=False)
    if kind == "b":
        return "Y" if v else "N"
    return v


def _value(v: Any, kind: str, where: str) -> Any:
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return [] if kind == "l" else None
    try:
        if kind == "j":
            return json.loads(v) if isinstance(v, str) else v
        if kind == "l":
            s = str(v).strip()
            if s.startswith("["):
                return json.loads(s)
            return [x.strip() for x in s.split(",") if x.strip()]
        if kind == "b":
            return str(v).strip().lower() in ("y", "yes", "true", "1", "예", "o")
        if kind == "i":
            return int(float(v))
        if kind == "f":
            return float(v)
    except (ValueError, TypeError) as exc:
        raise BundleError(f"{where}: 값을 읽을 수 없습니다 — {exc}") from exc
    return str(v) if not isinstance(v, str) else v


def write_bundle(b: dict[str, Any]) -> bytes:
    """묶음 → 엑셀(시트별 표 · 첫 시트 안내)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    flat = _flatten(b)
    wb = Workbook()
    ws = wb.active
    ws.title = "안내"
    ws.append(["항목", "값"])
    for k, v in (("형식", FORMAT), ("이름", b.get("name") or ""), ("설명", b.get("description") or ""),
                 ("올리는 곳", "관리자 › 데이터와 지식 › 업무 구성 — 미리보기에서 자료 연결 · 오류를 확인한 뒤 반영"),
                 ("자료", "지표 · 화면 연결은 '자료' 시트의 이름으로 원천을 가리킨다. 그 자료(필요한 열이 모두 있는 데이터셋)가 올라오면 자동으로 이어진다")):
        ws.append([k, v])
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E79")
    for sheet, key, cols in SHEETS:
        rows = flat.get(key) or []
        w = wb.create_sheet(sheet)
        w.append([c[0] for c in cols])
        for c in w[1]:
            c.font, c.fill = head, fill
        for r in rows:
            w.append([_cell(r.get(f), kind) for _, f, kind in cols])
        for i, (h, _, kind) in enumerate(cols, start=1):
            w.column_dimensions[w.cell(row=1, column=i).column_letter].width = 48 if kind == "j" else max(12, min(36, len(h) * 2 + 4))
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=False)
        w.freeze_panes = "A2"
    for c in ws[1]:
        c.font, c.fill = head, fill
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 14, 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _flatten(b: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    out = {k: list(b.get(k) or []) for _, k, _ in SHEETS if not k.startswith("_")}
    nodes, trig = [], []
    for w in b.get("workflows") or []:
        for i, n in enumerate(w.get("nodes") or [], start=1):
            nodes.append({"workflow": w["code"], "seq": n.get("seq") or i * 10, **n})
        for t in w.get("triggers") or []:
            trig.append({"workflow": w["code"], **t})
    out["_nodes"], out["_triggers"] = nodes, trig
    arts = []
    for a in b.get("artifacts") or []:
        s = json.dumps(a["payload"], ensure_ascii=False)
        for i in range(0, max(1, len(s)), CHUNK):
            arts.append({"key": a["key"], "part": i // CHUNK + 1, "chunk": s[i:i + CHUNK]})
    out["_artifacts"] = arts
    on = b.get("ontology") or {}
    rows = [{"kind": "class", "name": c.get("label"), "body": c} for c in on.get("classes") or []]
    rows += [{"kind": "relation", "name": r.get("type"), "body": r} for r in on.get("relations") or []]
    rows += [{"kind": "alias", "name": k, "body": v} for k, v in (on.get("aliases") or {}).items()]
    rows += [{"kind": "interface", "name": i.get("name"), "body": i} for i in on.get("interfaces") or []]
    out["_ontology"] = rows
    return out


def read_bundle(filename: str, content: bytes) -> dict[str, Any]:
    """엑셀 · JSON → 묶음. 형식 오류는 BundleError(시트 · 줄 · 칸)."""
    if len(content) > MAX_BYTES:
        raise BundleError(f"파일이 너무 큽니다(최대 {MAX_BYTES // 1024 // 1024}MB)")
    name = (filename or "").lower()
    if name.endswith(".json"):
        try:
            b = json.loads(content.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise BundleError(f"JSON 을 읽을 수 없습니다 — {exc}") from exc
        if not isinstance(b, dict):
            raise BundleError("JSON 최상위는 객체여야 합니다")
        return _normalize(b)
    if not name.endswith((".xlsx", ".xlsm")):
        raise BundleError("업무 구성 파일은 엑셀(.xlsx) 또는 JSON 입니다")
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise BundleError(f"엑셀을 읽을 수 없습니다 — {exc}") from exc
    b: dict[str, Any] = {}
    if "안내" in wb.sheetnames:
        for r in wb["안내"].iter_rows(min_row=2, values_only=True):
            if r and r[0] in ("이름", "설명", "형식"):
                b[{"이름": "name", "설명": "description", "형식": "format"}[r[0]]] = r[1] if len(r) > 1 else None
    flat: dict[str, list[dict[str, Any]]] = {}
    for sheet, key, cols in SHEETS:
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            continue
        head = [str(h).strip() if h is not None else "" for h in rows[0]]
        idx = {}
        for h, f, kind in cols:
            if h in head:
                idx[f] = (head.index(h), kind, h)
        items = []
        for n, r in enumerate(rows[1:], start=2):
            if not r or all(v is None or str(v).strip() == "" for v in r):
                continue
            items.append({f: _value(r[i] if i < len(r) else None, kind, f"[{sheet}] {n}줄 '{h}'") for f, (i, kind, h) in idx.items()})
        flat[key] = items
    for _, key, _ in SHEETS:
        if not key.startswith("_") and key in flat:
            b[key] = flat[key]
    wfs = {w["code"]: {**w, "nodes": [], "triggers": []} for w in b.get("workflows") or [] if w.get("code")}
    for n in sorted(flat.get("_nodes") or [], key=lambda x: (x.get("workflow") or "", x.get("seq") or 0)):
        w = wfs.get(n.pop("workflow", None) or "")
        if w is None:
            raise BundleError(f"[흐름단계] 업무흐름 시트에 없는 흐름: {n.get('key')}")
        w["nodes"].append(n)
    for t in flat.get("_triggers") or []:
        w = wfs.get(t.pop("workflow", None) or "")
        if w is None:
            raise BundleError(f"[흐름조건] 업무흐름 시트에 없는 흐름: {t.get('target')}")
        w["triggers"].append(t)
    if wfs:
        b["workflows"] = list(wfs.values())
    parts: dict[str, list[tuple[int, str]]] = {}
    for a in flat.get("_artifacts") or []:
        parts.setdefault(a["key"], []).append((a.get("part") or 1, a.get("chunk") or ""))
    arts = []
    for k, ps in parts.items():
        s = "".join(c for _, c in sorted(ps))
        try:
            arts.append({"key": k, "payload": json.loads(s)})
        except json.JSONDecodeError as exc:
            raise BundleError(f"[화면틀] '{k}' 내용을 JSON 으로 읽을 수 없습니다 — {exc}") from exc
    if arts:
        b["artifacts"] = arts
    on: dict[str, Any] = {"classes": [], "relations": [], "aliases": {}, "interfaces": []}
    for r in flat.get("_ontology") or []:
        kind = r.get("kind")
        if kind == "class":
            on["classes"].append(r["body"])
        elif kind == "relation":
            on["relations"].append(r["body"])
        elif kind == "alias":
            on["aliases"][r["name"]] = r["body"]
        elif kind == "interface":
            on["interfaces"].append(r["body"])
    if any(on.values()):
        b["ontology"] = on
    return _normalize(b)


def _normalize(b: dict[str, Any]) -> dict[str, Any]:
    if b.get("format") and b["format"] != FORMAT:
        raise BundleError(f"알 수 없는 형식: {b['format']} (지원: {FORMAT})")
    name = str(b.get("name") or "").strip()
    if not name:
        raise BundleError("구성 이름이 없습니다 — [안내] 시트 '이름' 또는 JSON name")
    out = {"format": FORMAT, "name": name[:150], "description": str(b.get("description") or "")}
    for _, key, _ in SHEETS:
        if not key.startswith("_"):
            out[key] = [x for x in b.get(key) or [] if isinstance(x, dict)]
    out["workflows"] = [x for x in b.get("workflows") or [] if isinstance(x, dict)]
    out["artifacts"] = [x for x in b.get("artifacts") or [] if isinstance(x, dict) and x.get("key")]
    on = b.get("ontology") or {}
    out["ontology"] = {"classes": list(on.get("classes") or []), "relations": list(on.get("relations") or []), "aliases": dict(on.get("aliases") or {}), "interfaces": list(on.get("interfaces") or [])}
    return out
