"""데이터셋 적재 — 파싱 · 품질 검증 · 버전 확정 · 롤백 (ING-01~07).

데이터 형식이 바뀌면 코드가 아니라 `cfg.dataset` 정의(컬럼 · 별칭 · 파서 · 품질 규칙)를 바꾼다.
당행(bank)과 외부(ext)는 정의의 origin 에 따라 다른 스키마에 저장된다.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from sqlalchemy import select, update
from sqlalchemy.orm import Session

from nhdip.config import get_settings
from nhdip.events import audit, emit
from nhdip.models import Dataset, source, utcnow


class IngestError(ValueError):
    pass


# ═════════════════════════ 파싱 ═════════════════════════
def _norm(s: Any) -> str:
    return re.sub(r"[\s_\-()·./]+", "", str(s or "")).lower()


_NOTE_RE = re.compile(r"^\s*(?:※|주\s*\)|\*|출처\s*[:：]|자료\s*[:：]|비고\s*[:：])")


def header_candidates(ds: Dataset) -> set[str]:
    """정의 컬럼의 이름 · 라벨 · 별칭(정규화) — 머리글 줄 찾기용."""
    return {_norm(c) for col in ds.columns for c in [col["name"], col.get("label", ""), *col.get("aliases", [])] if c}


def _header_index(rows: list[Any], cands: set[str] | None, max_scan: int = 20) -> int:
    """제목 · 기준 · 단위 줄 아래의 머리글 줄 — 정의 컬럼과 가장 많이(2개 이상) 맞는 첫 줄. 없으면 첫 줄(RD-09)."""
    if not cands:
        return 0
    best, best_n = 0, 0
    for i, r in enumerate(rows[:max_scan]):
        n = sum(1 for v in (r or []) if v is not None and _norm(v) in cands)
        if n > best_n:
            best, best_n = i, n
    return best if best_n >= 2 else 0


def _is_note(cells: Any) -> bool:
    """빈 줄 또는 주석 줄(첫 칸만 '※ · 주) · * · 출처:' 로 시작) — 행으로 읽지 않는다."""
    vals = [v for v in cells if v is not None and str(v).strip() != ""]
    return not vals or (len(vals) == 1 and bool(_NOTE_RE.match(str(vals[0]))))


def _read_table(content: bytes, filename: str, parser: dict[str, Any], cands: set[str] | None = None) -> list[dict[str, Any]]:
    name = filename.lower()
    if name.endswith(".json"):
        data = json.loads(content.decode("utf-8-sig"))
        if isinstance(data, dict):
            data = data.get("rows") or data.get("data") or []
        if not isinstance(data, list):
            raise IngestError("JSON 은 행 목록이어야 합니다")
        return [dict(r) for r in data]
    if name.endswith((".xlsx", ".xlsm")):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        ws = wb[parser["sheet"]] if parser.get("sheet") in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header_idx = int(parser["header_row"]) - 1 if parser.get("header_row") else _header_index(rows, cands)
        header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[header_idx])]
        out = []
        for r in rows[header_idx + 1 :]:
            if r is None or _is_note(r):
                continue
            out.append({header[i]: r[i] for i in range(min(len(header), len(r)))})
        return out
    if name.endswith((".csv", ".txt", ".tsv")):
        text = content.decode(parser.get("encoding", "utf-8-sig"), errors="replace")
        delim = "\t" if name.endswith(".tsv") else parser.get("delimiter", ",")
        skip = int(parser.get("skip_rows", 0))
        lines = text.splitlines()[skip:]
        if "skip_rows" not in parser and cands:
            lines = lines[_header_index(list(csv.reader(lines[:20], delimiter=delim)), cands) :]
        out = []
        for r in csv.DictReader(lines, delimiter=delim):
            r = {k: v for k, v in r.items() if k is not None}
            if not _is_note(r.values()):
                out.append(r)
        return out
    raise IngestError("지원 형식: .csv · .tsv · .xlsx · .json")


def _coerce(value: Any, typ: str) -> Any:
    if value is None or (isinstance(value, str) and value.strip() in ("", "-")):
        return None
    if typ in ("int", "integer"):
        try:
            return int(Decimal(str(value).replace(",", "")))
        except (InvalidOperation, ValueError) as exc:  # 'abc' 같은 값이 500 이 아니라 형 변환 오류로 모이도록
            raise ValueError(f"정수 아님: {value!r}") from exc
    if typ in ("num", "number", "decimal", "float"):
        try:
            return float(Decimal(str(value).replace(",", "")))
        except InvalidOperation as exc:
            raise ValueError(f"숫자 아님: {value!r}") from exc
    if typ == "bool":
        return str(value).strip().lower() in ("1", "y", "yes", "true", "t", "예")
    if typ == "date":
        if isinstance(value, (date, datetime)):
            return value.strftime("%Y-%m-%d")
        s = str(value).strip()
        for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y%m%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
            except ValueError:
                continue
        raise ValueError(f"날짜 아님: {value!r}")
    s = str(value).strip()
    if isinstance(value, float) and value.is_integer():
        s = str(int(value))
    return s


NO_COLUMNS_MSG = "컬럼 정의가 없습니다 — 관리자 [데이터셋]에서 컬럼을 정의하거나, 파일에서 컬럼을 추론해 정의하세요(정의 없이는 읽을 열을 알 수 없습니다)"
_AUTO_HEADER = re.compile(r"^col\d+$")  # 엑셀 머리글 칸이 비어 있을 때 읽기 단계에서 붙인 이름


def header_map(ds: Dataset, headers: list[str]) -> tuple[dict[str, str], list[str]]:
    """정의 컬럼 → 원본 헤더 매핑(이름 · 라벨 · 별칭을 정규화해 비교)과 필수인데 못 찾은 열의 오류 메시지."""
    hmap: dict[str, str] = {}
    errors: list[str] = []
    for col in ds.columns:
        cands = [col["name"], col.get("label", ""), *col.get("aliases", [])]
        found = next((h for h in headers if _norm(h) in {_norm(c) for c in cands if c}), None)
        if found is not None:
            hmap[col["name"]] = found
        elif col.get("required"):
            errors.append(f"필수 컬럼 없음: {col['name']} (허용 헤더: {', '.join(c for c in cands if c)})")
    return hmap, errors


def unmapped_headers(ds: Dataset, raw_rows: list[dict[str, Any]]) -> list[str]:
    """파일에는 있지만 정의에 없어 저장하지 않는 열 머리글 — 조용히 버려지지 않도록 검사 결과에 알린다."""
    if not raw_rows or not ds.columns:
        return []
    headers = [str(h) for h in raw_rows[0].keys()]
    used = set(header_map(ds, headers)[0].values())
    return [h for h in headers if h not in used and h.strip() and not _AUTO_HEADER.match(h)]


def map_rows(ds: Dataset, raw_rows: list[dict[str, Any]], warnings: list[str] | None = None) -> tuple[list[dict[str, Any]], list[str]]:
    """원본 헤더 → 정의 컬럼 매핑 + 형 변환. (행, 오류 목록)
    자동으로 더한 열(`auto`, M.24)은 형 변환이 안 되는 값을 비워 두고 warnings 에만 적는다 — 처음 본 파일로 정한 타입 때문에 다음 파일이 막히지 않게."""
    if not ds.columns:
        return [], [NO_COLUMNS_MSG]  # 정의가 비면 모든 행이 빈 값으로 '통과' 저장되던 것을 막는다
    if not raw_rows:
        return [], ["행이 없습니다"]
    headers = list(raw_rows[0].keys())
    hmap, errors = header_map(ds, headers)
    rows: list[dict[str, Any]] = []
    for i, raw in enumerate(raw_rows, start=1):
        row: dict[str, Any] = {}
        for col in ds.columns:
            h = hmap.get(col["name"])
            try:
                row[col["name"]] = _coerce(raw.get(h) if h else None, col.get("type", "str"))
            except ValueError as exc:
                if col.get("auto") and not col.get("required"):
                    if warnings is not None and len(warnings) < 50:
                        warnings.append(f"{i}행 {col['name']}: {exc}")
                else:
                    errors.append(f"{i}행 {col['name']}: {exc}")
                row[col["name"]] = None
        rows.append(row)
    return rows, errors


# ═════════════════════════ 정의에 없는 열 — 자동으로 더하기 (M.24 · DOB-17 · Q212) ═════════════════════════
AUTO_COLUMNS_SETTING = "dataset.auto_columns"
_VIEW_FIELDS = ("code", "name", "origin", "kind", "description", "file_pattern", "refresh_cycle", "key_columns", "as_of_column", "mapping", "quality_rules", "parser", "materialize", "active")


def auto_columns_on(session: Session | None) -> bool:
    """정의에 없는 열을 정의에 더해 저장할지 — 설정 dataset.auto_columns(기본 true). false 면 예전처럼 안내만 하고 저장하지 않는다."""
    if session is None:
        return True
    from nhdip.models import Setting

    st = session.get(Setting, AUTO_COLUMNS_SETTING)
    return not (st is not None and st.value is False)


def new_columns_from(ds: Any, raw_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """파일에는 있지만 정의에 없는 열 → 정의에 더할 열(선택 · `auto` 표시). 이름은 머리글에서 만들고 기존 이름과 정규화 비교로 겹치지 않게,
    타입은 값으로 추론(섞이면 문자), 라벨 = 원래 머리글(다음 업로드부터 그 머리글로 맞춘다)."""
    from .infer import MAX_COLUMNS, SCAN_ROWS, _column_type, _slug

    heads = unmapped_headers(ds, raw_rows)
    if not heads:
        return []
    used = {str(c.get("name")).lower() for c in ds.columns}
    normed = {_norm(c.get("name")) for c in ds.columns}
    room = MAX_COLUMNS - len(ds.columns)
    sample = raw_rows[:SCAN_ROWS]
    out: list[dict[str, Any]] = []
    for i, h in enumerate(heads):
        if len(out) >= room:
            break
        name = _slug(h, used, len(ds.columns) + i)
        n = 2
        while _norm(name) in normed or not _norm(name):
            name = f"{name}_{n}"[:60]
            n += 1
        normed.add(_norm(name))
        typ, _is_ym, _note = _column_type(h, [r.get(h) for r in sample])
        out.append({"name": name[:60], "label": str(h).strip()[:200], "type": typ, "required": False, "aliases": [], "auto": True})
    return out


def with_new_columns(session: Session | None, ds: Dataset, raw_rows: list[dict[str, Any]], source: dict[str, Any] | None) -> tuple[Any, list[dict[str, Any]]]:
    """(읽을 때 쓸 정의, 더할 열). 원본 보고서 양식(어댑터)은 열이 정해져 있어 더하지 않는다. 정의 점검을 통과하지 못하면 더하지 않는다(안내만).
    정의 자체는 바꾸지 않는다 — 확정된 업로드에서만 `extend_definition` 이 바꾼다."""
    from types import SimpleNamespace

    from .infer import validate_definition

    if (source or {}).get("adapter") or not ds.columns or not auto_columns_on(session):
        return ds, []  # 예전 원본 양식 어댑터는 열이 그 정의에 맞춰 정해져 있다 — 보고서 양식 읽기(layout)는 새 열을 더한다
    new = new_columns_from(ds, raw_rows)
    if not new:
        return ds, []
    cols = [*ds.columns, *new]
    if validate_definition({"code": ds.code, "name": ds.name, "origin": ds.origin, "columns": cols, "key_columns": ds.key_columns, "as_of_column": ds.as_of_column}):
        return ds, []
    return SimpleNamespace(**{k: getattr(ds, k, None) for k in _VIEW_FIELDS}, columns=cols), new


def extend_definition(session: Session, ds: Dataset, new: list[dict[str, Any]], actor: str | None, file_name: str | None) -> None:
    """확정된 업로드가 가져온 새 열을 정의에 더한다 — 정의 변경 이벤트로 구조 파악(온톨로지 구축) · 설명 카드 · 그래프가 새 열을 본다."""
    have = {_norm(c.get("name")) for c in ds.columns}
    add = [c for c in new if _norm(c["name"]) not in have]
    if not add:
        return
    ds.columns = [*ds.columns, *add]  # 새 목록으로 바꿔 넣는다(JSONB 변경 감지)
    emit(session, "config.changed", {"kind": "dataset", "code": ds.code, "columnsAdded": [c["name"] for c in add]}, actor or "system")
    audit(session, actor, "dataset.columns_auto_added", "dataset", ds.code, {"columns": [{"name": c["name"], "label": c["label"], "type": c["type"]} for c in add], "file": file_name})


# ═════════════════════════ 품질 검증 ═════════════════════════
def _ym_index(v: Any) -> int | None:
    """'202608' · '2026-08' · '2026-08-31' → 월 일련번호."""
    t = re.sub(r"[^0-9]", "", str(v or ""))
    if len(t) < 6:
        return None
    y, m = int(t[:4]), int(t[4:6])
    return y * 12 + (m - 1) if 1 <= m <= 12 else None


def quality_check(ds: Dataset, rows: list[dict[str, Any]], parse_errors: list[str], previous: list[dict[str, Any]] | None = None, privacy: dict[str, Any] | None = None, unmapped: list[str] | None = None,
                  added: list[dict[str, Any]] | None = None, warnings: list[str] | None = None) -> dict[str, Any]:
    checks: list[dict[str, Any]] = []

    def add(rule: str, passed: bool, detail: str, severity: str = "error"):
        checks.append({"rule": rule, "passed": passed, "detail": detail, "severity": severity})

    add("parse", not parse_errors, "; ".join(parse_errors[:10]) or "헤더·형 변환 정상")
    if added:
        # 정의에 없던 열 — 정의에 더해 값까지 저장한다(M.24 · DOB-17). 통과 · 실패에 영향 없음
        shown = ", ".join(f"{c['label']}({c['type']})" for c in added[:8]) + (f" 외 {len(added) - 8}개" if len(added) > 8 else "")
        add("columns_added", True, f"정의에 없던 열 {len(added)}개를 정의에 더해 저장합니다: {shown}", "info")
    if warnings:
        add("auto_column_values", True, f"자동으로 더한 열에서 타입이 맞지 않는 값 {len(warnings)}건은 비워 두었습니다: " + "; ".join(warnings[:5]), "warning")
    if unmapped:
        # 통과·실패에 영향 없는 안내 — 정의에 없는 열은 저장하지 않는다는 것을 알려, 필요하면 정의에 열을 더하게 한다
        shown = ", ".join(unmapped[:8]) + (f" 외 {len(unmapped) - 8}개" if len(unmapped) > 8 else "")
        add("unmapped_headers", True, f"정의에 없는 열 {len(unmapped)}개는 저장하지 않았습니다: {shown} — 관리자 [데이터셋]의 '파일에서 새 열 찾기'로 정의에 더할 수 있습니다", "info")
    for col in ds.columns:
        if col.get("required"):
            missing = sum(1 for r in rows if r.get(col["name"]) is None)
            add(f"not_null:{col['name']}", missing == 0, f"결측 {missing}건")
    if ds.key_columns:
        keys = [tuple(r.get(k) for k in ds.key_columns) for r in rows]
        dup = len(keys) - len(set(keys))
        add("unique_key", dup == 0, f"키({', '.join(ds.key_columns)}) 중복 {dup}건")
    for rule in ds.quality_rules or []:
        kind = rule.get("rule")
        sev = rule.get("severity", "error")
        col = rule.get("column")
        if kind == "range":
            lo, hi = rule.get("min"), rule.get("max")
            bad = [r[col] for r in rows if r.get(col) is not None and ((lo is not None and r[col] < lo) or (hi is not None and r[col] > hi))]
            add(f"range:{col}", not bad, f"범위({lo}~{hi}) 밖 {len(bad)}건", sev)
        elif kind == "allowed":
            allowed = set(rule.get("values", []))
            bad = {r.get(col) for r in rows if r.get(col) not in allowed}
            add(f"allowed:{col}", not bad, f"허용 외 값 {sorted(map(str, bad))[:5]}", sev)
        elif kind == "le_column":
            other = rule.get("other")
            bad = sum(1 for r in rows if r.get(col) is not None and r.get(other) is not None and r[col] > r[other])
            add(f"le:{col}<={other}", bad == 0, f"{col} > {other} 인 행 {bad}건", sev)
        elif kind == "sum_equals":
            # 상세 행 합계 = 합계 행 값 (허용 오차 %) — 그룹 컬럼별
            flag, val, group, tol = rule["total_flag"], rule["column"], rule.get("group"), float(rule.get("tolerance_pct", 0.1))
            groups = {r.get(group) for r in rows} if group else {None}
            worst = 0.0
            for g in groups:
                part = [r for r in rows if not group or r.get(group) == g]
                total = sum(r[val] or 0 for r in part if r.get(flag))
                detail = sum(r[val] or 0 for r in part if not r.get(flag))
                if total:
                    worst = max(worst, abs(detail - total) / abs(total) * 100)
            add(f"sum_equals:{val}", worst <= tol, f"상세 합계 오차 최대 {worst:.3f}% (허용 {tol}%)", sev)
        elif kind == "period_continuity":
            # 기간 연속성 — 그룹별로 처음~마지막 월 사이 빠진 월 (1.4)
            group = rule.get("group")
            gaps: list[str] = []
            for g in sorted({r.get(group) for r in rows} if group else {None}, key=str):
                idx = sorted({i for r in rows if (not group or r.get(group) == g) for i in [_ym_index(r.get(col))] if i is not None})
                missing = [i for i in range(idx[0], idx[-1] + 1) if i not in set(idx)] if idx else []
                if missing:
                    gaps.append(f"{g or '전체'}: " + ", ".join(f"{i // 12}{i % 12 + 1:02d}" for i in missing[:6]))
            add(f"continuity:{col}", not gaps, "; ".join(gaps[:5]) or "빠진 기간 없음", sev)
        elif kind == "max_change":
            # 급변 경고 — 직전 현재 버전 대비 같은 키 값의 변화율이 기준(%)을 넘는 행 (1.4)
            if previous is None:
                continue
            keys, pct = rule.get("keys") or ds.key_columns, float(rule.get("pct", 50))
            before = {tuple(r.get(k) for k in keys): r.get(col) for r in previous}
            spikes = []
            for r in rows:
                b, a = before.get(tuple(r.get(k) for k in keys)), r.get(col)
                if isinstance(a, (int, float)) and isinstance(b, (int, float)) and b != 0 and abs(a - b) / abs(b) * 100 > pct:
                    spikes.append(f"{'/'.join(str(r.get(k)) for k in keys)} {b:,.0f}→{a:,.0f}")
            add(f"max_change:{col}", not spikes, f"직전 버전 대비 {pct:g}% 넘게 바뀐 행 {len(spikes)}건" + (f" — {'; '.join(spikes[:3])}" if spikes else ""), sev)
        elif kind == "min_rows":
            add("min_rows", len(rows) >= int(rule.get("value", 1)), f"{len(rows)}행 (최소 {rule.get('value')})", sev)
    # 개인정보 형식 값 검사 (B.32) — 주민등록번호 · 휴대전화 · 이메일 · 카드번호가 섞이면 적재 거부 (설정 privacy.policy)
    from nhdip.privacy import DEFAULTS as PRIVACY_DEFAULTS, LABELS as PII_LABELS, scan_rows

    pol = {**PRIVACY_DEFAULTS, **(privacy or {})}
    if pol.get("pii_scan", True):
        allow = set(pol.get("pii_allow_columns") or []) | {c["name"] for c in ds.columns if c.get("pii_ok")}
        res = scan_rows(rows, allow)
        found = ", ".join(f"{PII_LABELS[k]} {n}건" for k, n in res["found"].items())
        where = "; ".join(f"{x['row']}행 {x['column']}({x['kind']})" for x in res["samples"][:5])
        add("pii", not res["found"], f"개인정보 형식 값 {found} — {where} (허용 열은 설정 privacy.policy.pii_allow_columns)" if res["found"] else "개인정보 형식 값 없음")
    failed = [c for c in checks if not c["passed"] and c["severity"] == "error"]
    warned = [c for c in checks if not c["passed"] and c["severity"] != "error"]
    status = "failed" if failed else ("warning" if warned else "passed")
    return {"status": status, "checks": checks, "rows": len(rows)}


# ═════════════════════════ 버전 확정 ═════════════════════════
def _store_file(origin: str, code: str, filename: str, content: bytes) -> tuple[str, str]:
    sha = hashlib.sha256(content).hexdigest()
    base = os.path.join(get_settings().filestore_dir, origin, code)
    try:
        os.makedirs(base, exist_ok=True)
        path = os.path.join(base, f"{sha[:16]}_{os.path.basename(filename)}")
        if not os.path.exists(path):
            with open(path, "wb") as fh:
                fh.write(content)
    except OSError:
        path = ""  # 파일 볼륨이 없으면(테스트 등) 체크섬만 남긴다
    return sha, path


def to_date(v: Any) -> str | None:
    """기준일 컬럼 값(YYYYMM · YYYY-MM-DD · YYYY.MM) → ISO 날짜 문자열(월 단위면 말일)."""
    import calendar

    if v is None:
        return None
    s = str(v).strip().replace(".", "").replace("/", "")
    if len(s) == 6 and s.isdigit():
        y, m = int(s[:4]), int(s[4:])
        if 1 <= m <= 12:
            return f"{y:04d}-{m:02d}-{calendar.monthrange(y, m)[1]:02d}"
    s2 = str(v).strip()
    if len(s2) == 10 and s2[4] == "-":
        return s2
    if len(s) == 8 and s.isdigit():
        return f"{s[:4]}-{s[4:6]}-{s[6:]}"
    return None


def _row_key(ds: Dataset, row: dict[str, Any]) -> str | None:
    if not ds.key_columns:
        return None
    return "|".join(str(row.get(k)) for k in ds.key_columns)[:300]


def ingest_rows(
    session: Session,
    ds: Dataset,
    rows: list[dict[str, Any]],
    *,
    quality: dict[str, Any],
    file_name: str | None,
    file_sha: str | None = None,
    file_path: str | None = None,
    as_of: str | None = None,
    data_origin: str = "simulated",
    actor: str | None = None,
    note: str | None = None,
    force: bool = False,
):
    """버전 생성 → (품질 통과 시) 확정 + 현재 버전 교체 + 아웃박스 이벤트 — 한 트랜잭션."""
    from nhdip.locks import lock

    t = source(ds.origin)
    lock(session, f"ds:{ds.origin}:{ds.code}")  # 동시 업로드 · 되돌리기 직렬화 — 버전 번호 · 현재 버전 교체가 겹치지 않도록 (B.32)
    last = session.execute(
        select(t.version.version).where(t.version.dataset_code == ds.code).order_by(t.version.version.desc()).limit(1)
    ).scalar()
    commit = quality["status"] != "failed" or force
    if not as_of and ds.as_of_column:
        # 버전 기준일 = 파일 안 마지막 기간(Q141) — 여러 달이 든 파일(2020 원본 1~12월 · 월별 부도율)이 서로를 밀어내지 않게
        vals = {to_date(r.get(ds.as_of_column)) for r in rows if r.get(ds.as_of_column) is not None} - {None}
        if vals:
            as_of = max(vals)
    v = t.version(
        dataset_code=ds.code,
        version=(last or 0) + 1,
        as_of=date.fromisoformat(as_of) if as_of else None,
        data_origin=data_origin,
        status="committed" if commit else "rejected",
        is_current=False,
        file_name=file_name,
        file_sha256=file_sha,
        file_path=file_path,
        row_count=len(rows),
        quality=quality,
        note=note,
        uploaded_by=actor,
        committed_at=utcnow() if commit else None,
    )
    session.add(v)
    session.flush()
    for i, r in enumerate(rows, start=1):
        row_as_of = to_date(r.get(ds.as_of_column)) if ds.as_of_column else None
        session.add(
            t.row(
                version_id=v.id,
                dataset_code=ds.code,
                row_no=i,
                row_key=_row_key(ds, r),
                as_of=date.fromisoformat(row_as_of) if row_as_of else None,
                data=r,
            )
        )
    if commit:
        # 같은 기준일 구획(partition)의 이전 현재 버전만 교체 — 다른 기준일 데이터는 유지
        same = t.version.as_of == v.as_of if v.as_of else t.version.as_of.is_(None)
        session.execute(update(t.version).where(t.version.dataset_code == ds.code, t.version.id != v.id, same).values(is_current=False))
        v.is_current = True
        emit(
            session,
            "dataset.committed",
            {"dataset": ds.code, "origin": ds.origin, "kind": ds.kind, "version": v.version, "version_id": v.id, "rows": len(rows)},
            actor,
        )
    audit(session, actor, "dataset.upload", "dataset", ds.code, {"version": v.version, "status": v.status, "quality": quality["status"], "file": file_name})
    return v


def _read_any(content: bytes, filename: str, ds: Dataset) -> tuple[list[dict[str, Any]], dict[str, Any] | None]:
    """파일 → 원본 행. 읽는 방법은 정의의 parser 가 정한다(파일에서 정의를 만들 때 정해짐):
      layout=report   업무 보고서 양식(다단 머리글 · 계층 · 월 열) — 파일 모양에서 구조를 만든다(data/layout.py, 정한 양식 · 코드 없음 · M.25)
      adapter=…       예전 정의가 명시한 원본 양식 어댑터(2633 · 2020 · 7790 — 그 정의에 맞춘 열로 변환). 명시하지 않으면 쓰지 않는다
      그 밖           일반 표(머리글 1줄 · 제목 줄 건너뜀)
    보고서 양식으로 정한 정의에 일반 표가 오면(또는 그 반대) 파일 모양을 따른다."""
    parser = ds.parser or {}
    if filename.lower().endswith((".xlsx", ".xlsm")):
        if parser.get("adapter") not in (None, "", "none"):
            from .adapters import AdapterError, parse_report

            try:
                parsed = parse_report(content, filename, ds.code, parser)
            except AdapterError as exc:
                raise IngestError(str(exc)) from exc
            if parsed is not None:
                if not parsed.rows:
                    raise IngestError(f"{parsed.adapter}: 읽은 값이 없습니다")
                return parsed.rows, {"adapter": parsed.adapter, "meta": parsed.meta, "notes": parsed.notes}
        from .layout import read_report

        rep = read_report(content, filename)
        if rep is not None:
            return rep.rows, {"layout": "report", **{k: v for k, v in rep.info.items() if k != "layout"}}
    return _read_table(content, filename, parser, header_candidates(ds)), None


def _read_mapped(session: Session | None, ds: Dataset, filename: str, content: bytes) -> tuple[list[dict[str, Any]], dict[str, Any] | None, Any, list[dict[str, Any]], list[dict[str, Any]], list[str], list[str]]:
    """(원본 행, 어댑터 정보, 읽은 정의(새 열 포함), 더할 열, 행, 오류, 경고) — 확정 · 미리보기 · 요청 검사가 같은 규칙으로 읽는다."""
    raw, source = _read_any(content, filename, ds)
    view, added = with_new_columns(session, ds, raw, source)
    warns: list[str] = []
    rows, errors = map_rows(view, raw, warns)
    return raw, source, view, added, rows, errors, warns


def ingest_file(session: Session, ds: Dataset, filename: str, content: bytes, **kw):
    raw, source, view, added, rows, errors, warns = _read_mapped(session, ds, filename, content)
    from nhdip.privacy import policy as privacy_policy

    _, prev = current_rows(session, ds)
    quality = quality_check(view, rows, errors, previous=prev, privacy=privacy_policy(session), unmapped=unmapped_headers(view, raw), added=added, warnings=warns)
    if source:
        quality["source"] = source  # 원본 양식 · 머리글 메타(조회일자 · 조회사무소 · 기준일 · 단위) · 변환 메모
    sha, path = _store_file(ds.origin, ds.code, filename, content)
    v = ingest_rows(session, ds, rows, quality=quality, file_name=filename, file_sha=sha, file_path=path, **kw)
    if added and v.status == "committed":
        extend_definition(session, ds, added, kw.get("actor"), filename)
    return v


def preview_file(ds: Dataset, filename: str, content: bytes, session: Session | None = None) -> dict[str, Any]:
    raw, source, view, added, rows, errors, warns = _read_mapped(session, ds, filename, content)
    return {"headers": list(raw[0].keys()) if raw else [], "rows": rows[:20], "quality": quality_check(view, rows, errors, unmapped=unmapped_headers(view, raw), added=added, warnings=warns), "source": source,
            "columnsAdded": added}


def check_file(session: Session, ds: Dataset, filename: str, content: bytes, sample: int = 10) -> dict[str, Any]:
    """업무 사용자 자료 요청의 사전 검사(Phase C · UPL-01) — 확정과 같은 품질 검사(직전 현재 버전 대비 · 개인정보 정책), 저장 없음.
    정의에 없는 열은 반영(승인) 때 정의에 더해 저장된다 — 여기서는 무엇이 더해지는지만 알린다(M.24)."""
    from nhdip.privacy import policy as privacy_policy

    raw, source, view, added, rows, errors, warns = _read_mapped(session, ds, filename, content)
    _, prev = current_rows(session, ds)
    quality = quality_check(view, rows, errors, previous=prev, privacy=privacy_policy(session), unmapped=unmapped_headers(view, raw), added=added, warnings=warns)
    if source:
        quality["source"] = source
    return {"headers": list(raw[0].keys()) if raw else [], "columns": [c["name"] for c in view.columns], "sample": rows[:sample], "rows": len(rows), "quality": quality, "columnsAdded": added}


def rollback(session: Session, ds: Dataset, version_id: int, actor: str | None):
    """같은 기준일 구획 안에서 지정 버전을 현재 버전으로 되돌린다(이후 버전은 rolled_back 표시)."""
    from nhdip.locks import lock

    t = source(ds.origin)
    lock(session, f"ds:{ds.origin}:{ds.code}")
    target = session.get(t.version, version_id)
    if not target or target.dataset_code != ds.code or target.status not in ("committed", "rolled_back"):
        raise IngestError("되돌릴 수 있는 버전이 아닙니다")
    same = t.version.as_of == target.as_of if target.as_of else t.version.as_of.is_(None)
    for cur in session.execute(select(t.version).where(t.version.dataset_code == ds.code, t.version.is_current, same)).scalars():
        cur.is_current = False
        if cur.version > target.version:
            cur.status = "rolled_back"
    target.is_current = True
    target.status = "committed"
    target.committed_at = utcnow()
    emit(session, "dataset.committed", {"dataset": ds.code, "origin": ds.origin, "kind": ds.kind, "version": target.version, "version_id": target.id, "rollback": True}, actor)
    audit(session, actor, "dataset.rollback", "dataset", ds.code, {"to_version": target.version, "as_of": str(target.as_of) if target.as_of else None})
    return target


def latest_as_of(session: Session) -> str | None:
    """올라온 자료의 최신 기준일(현재 확정 버전 중 가장 늦은 날, ISO) — 없으면 None."""
    from sqlalchemy import func

    dates = [session.execute(select(func.max(source(o).version.as_of)).where(source(o).version.is_current)).scalar() for o in ("bank", "ext")]
    dates = [d for d in dates if d]
    return max(dates).isoformat() if dates else None


def display_as_of(session: Session) -> str | None:
    """화면 기준일 — 설정 ui.as_of 를 관리자가 정했으면 그 값, 비어 있으면 자료의 최신 기준일 (Phase I · 하드코딩 제거)."""
    from nhdip.models import Setting

    s = session.get(Setting, "ui.as_of")
    if s is not None and s.value:
        return str(s.value)
    return latest_as_of(session)


def current_versions(session: Session, ds: Dataset) -> list[Any]:
    t = source(ds.origin)
    return list(
        session.execute(
            select(t.version).where(t.version.dataset_code == ds.code, t.version.is_current).order_by(t.version.committed_at.desc(), t.version.id.desc())
        ).scalars()
    )


class VersionSet:
    """여러 기준일 구획의 현재 버전 묶음 — 계보 기록·표시용."""

    def __init__(self, versions: list[Any]):
        self.versions = versions
        latest = versions[0]
        self.id = latest.id
        self.version = latest.version if len(versions) == 1 else "+".join(str(v.version) for v in sorted(versions, key=lambda x: x.version))
        self.file_name = latest.file_name if len(versions) == 1 else f"{len(versions)}개 파일"
        self.data_origin = "actual" if all(v.data_origin == "actual" for v in versions) else "simulated"
        self.as_of = max((v.as_of for v in versions if v.as_of), default=None)
        self.committed_at = latest.committed_at


def current_rows(session: Session, ds: Dataset) -> tuple[VersionSet | None, list[dict[str, Any]]]:
    """모든 기준일 구획의 현재 버전 행. 같은 기준일 값이 여러 버전에 있으면 최근 확정 버전을 쓴다."""
    t = source(ds.origin)
    versions = current_versions(session, ds)
    if not versions:
        return None, []
    claimed: set[str] = set()
    out: list[dict[str, Any]] = []
    for v in versions:
        rows = session.execute(select(t.row.data).where(t.row.version_id == v.id).order_by(t.row.row_no)).scalars().all()
        if ds.as_of_column:
            mine = {str(r.get(ds.as_of_column)) for r in rows}
            rows = [r for r in rows if str(r.get(ds.as_of_column)) not in claimed]
            claimed |= mine
        # 행 단위 계보(어느 버전·파일에서 왔는지) — 계산 근거 표시에 사용
        out.extend({**r, "__v": v.version, "__file": v.file_name} for r in rows)
    return VersionSet(versions), out
