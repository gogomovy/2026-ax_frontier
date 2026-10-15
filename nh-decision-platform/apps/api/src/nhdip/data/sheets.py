"""엑셀 시트 읽기 공통 함수 — 값 · 들여쓰기 · 셀 위치 · 머리글 안내 줄(조회일자 · 기준일 · 단위 …).

특정 양식 · 데이터셋을 모른다. 보고서 모양 표 읽기(`data/layout.py`)가 쓴다.
안내 줄의 조회자 · 담당자 같은 개인 이름은 저장하지 않는다(DATA-06).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any

FILE_RE = re.compile(r"\((\d{6})\)\s*(\d{4})")
MONTH_RE = re.compile(r"^\s*(\d{1,2})\s*월\s*$")
UNIT_RE = re.compile(r"\[\s*([\d,]+)\s*\]\s*원")
UNIT_WORDS = {"원": 1, "천원": 1_000, "백만원": 1_000_000, "억원": 100_000_000, "조원": 1_000_000_000_000}


@dataclass
class Sheet:
    name: str
    cells: list[list[Any]]  # 값
    indents: list[list[int]]  # 들여쓰기 수준(셀 서식)

    def text(self, r: int, c: int) -> str:
        v = self.cells[r][c] if r < len(self.cells) and c < len(self.cells[r]) else None
        return "" if v is None else str(v)

    def ref(self, r: int, c: int) -> str:
        from openpyxl.utils import get_column_letter

        return f"{self.name}!{get_column_letter(c + 1)}{r + 1}"


def load_sheets(content: bytes) -> list[Sheet]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)  # 들여쓰기 서식을 읽기 위해 일반 모드
    out = []
    for ws in wb.worksheets:
        cells, indents = [], []
        for row in ws.iter_rows():
            cells.append([c.value for c in row])
            indents.append([int(getattr(c.alignment, "indent", 0) or 0) for c in row])
        out.append(Sheet(ws.title, cells, indents))
    return out


def file_info(filename: str) -> dict[str, str]:
    """파일 이름 '(YYYYMM)NNNN_…' 의 기준년월 · 보고서 번호."""
    m = FILE_RE.search(filename or "")
    return {"ym": m.group(1), "report_code": m.group(2)} if m else {}


def _level(sheet: Sheet, r: int, c: int) -> int:
    """들여쓰기 수준 = 셀 들여쓰기 + 앞 공백(반각 2칸 · 전각 1칸 = 1단계)."""
    raw = sheet.text(r, c)
    lead = len(raw) - len(raw.lstrip(" \u3000\t"))
    spaces = raw[:lead]
    return sheet.indents[r][c] + spaces.count("\u3000") + spaces.count("\t") + spaces.count(" ") // 2


def _num(v: Any) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    try:
        f = float(s.strip("()"))
    except ValueError:
        return None
    return -f if neg else f


def _header_meta(sheet: Sheet, upto: int) -> dict[str, Any]:
    """머리글 블록(헤더 행 위) — '키 : 값' · '키 : [값]' 조각을 모은다. 개인 이름 항목은 버린다."""
    meta: dict[str, Any] = {}
    skip = {"조회자", "담당자", "담당부서", "작성자"}
    for r in range(min(upto, len(sheet.cells))):
        line = " ".join(sheet.text(r, c) for c in range(len(sheet.cells[r])) if sheet.text(r, c).strip())
        if not line:
            continue
        um = UNIT_RE.search(line)
        if um:
            meta["unit_multiplier"] = int(um.group(1).replace(",", ""))
        else:
            uw = re.search(r"단위\s*[:：]?\s*(조원|억원|백만원|천원|원)", line)
            if uw:
                meta["unit_multiplier"] = UNIT_WORDS[uw.group(1)]
        for k, v in re.findall(r"([가-힣A-Za-z ]{2,10}?)\s*:\s*(\[[^\]]*\]|[^,:\[]+?)(?=,|\s+[가-힣A-Za-z ]{2,10}\s*:|$)", line):
            key = k.replace(" ", "")
            if key in skip or key.startswith("단위"):
                continue
            meta[key] = v.strip().strip("[]").strip()
        if "title" not in meta and r == 0 and ":" not in line:
            meta["title"] = line.strip()
    if "기준일" in meta:
        m = re.search(r"(\d{4})[-.](\d{2})", meta["기준일"])
        if m:
            meta["base_ym"] = m.group(1) + m.group(2)
    return meta
