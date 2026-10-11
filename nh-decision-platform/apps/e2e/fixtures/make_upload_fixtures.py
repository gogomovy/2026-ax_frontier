"""업로드 시험 파일 생성기 (B.40) — 정의서 · 데이터셋 정의에 맞춘 '다음 달(2026.09) 마감 자료' 엑셀 · PDF.

실제 사용자가 [데이터셋 · 문서] 화면에 올리는 파일을 흉내 낸다. 기준 값은 플랫폼에 확정된 2026.08 데이터(시드 시뮬레이션)이고,
여기에 현실적인 한 달 변화를 결정적 난수(seed 20260930)로 더한다. 모든 파일은 '시뮬레이션 시험 자료' 로 표시한다.

시나리오(2026.09):
  · 7790 상품별 충당금 — 포용금융 상품 충당금 +2~7%, 그 외 −1~+2% (포용금융 비중 상승 지속)
  · 신용평가모형 등급별 모니터링 — 관측 +0.3~1.2%, 소기업 6B 부도 196 → 205건(실제부도율 8.23%, 최대 PD 8.933% 이내)
  · Oblg_grd_def — 대출잔액 ±1%, 등급 이동 1.5%, 수원영통지점이 마케팅 후보 2개(신동에너지 50억 · 광교바이오팜 30억) 신규 유치
  · TB_JC_PM_C_LA(소매 익스포저) — 상환 0~1.2%, 연체 일부 변동, 수원영통지점 새희망홀씨 8건(25~35백만원) 신규
  · 업적평가 득점 현황 — 우량기업여신 · 서민금융 실적 반영(득점 = 실적 ÷ 목표 × 배점, 상한 배점)
  · PDF 2종 — 외부 정책(서민금융 공급 확대) · 내부 공문(소기업 6B 점검 요청)

실행(플랫폼 api 컨테이너 안 — DB 의 현재 확정 버전을 읽는다):
  docker compose exec -T api python - < apps/e2e/fixtures/make_upload_fixtures.py   # 출력: /tmp/upload-fixtures (컨테이너)
  OUT=/out python make_upload_fixtures.py
"""
from __future__ import annotations

import os
import random
from decimal import ROUND_HALF_UP, Decimal

import openpyxl
from openpyxl.styles import Alignment, Font

from nhdip.data.datasets import current_rows
from nhdip.db import SessionLocal
from nhdip.models import Dataset, MetricValue
from nhdip.models.runtime import CalcRun

OUT = os.environ.get("OUT", "/tmp/upload-fixtures")
YM, ASOF = "202609", "2026-09-30"
PREV_YM, PREV_ASOF = "202608", "2026-08-31"
OFFICE = "001068"
rng = random.Random(20260930)
os.makedirs(OUT, exist_ok=True)
s = SessionLocal()


def rows(code: str) -> list[dict]:
    _, r = current_rows(s, s.get(Dataset, code))
    return [{k: v for k, v in x.items() if not k.startswith("__")} for x in r]


def q(v, places="1") -> Decimal:
    return Decimal(str(v)).quantize(Decimal(places), rounding=ROUND_HALF_UP)


def won10k(v: Decimal) -> float:
    """만원 단위 반올림(원 단위 원장 금액처럼)."""
    return float((v / 10000).quantize(Decimal(1), rounding=ROUND_HALF_UP) * 10000)


def save(wb: openpyxl.Workbook, name: str) -> str:
    path = os.path.join(OUT, name)
    wb.save(path)
    return path


def table(name: str, sheet: str, header: list[str], data: list[list]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in data:
        ws.append(r)
    return save(wb, name)


made: list[tuple[str, str, str]] = []

# ── 1. 7790 상품별 충당금 (원본 양식: 머리글 · 다단 헤더 · X 열 총계 · 합계 행)
prov = [r for r in rows("PROVISION_BY_PRODUCT") if str(r["ym"]) == PREV_YM]
inc = {r["product_code"] for r in rows("INCLUSIVE_PRODUCTS")}
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "7790"
ws["A1"] = "7790 상품별 충당금 산출현황 요약 (시뮬레이션 시험 자료)"
ws["A1"].font = Font(bold=True, size=13)
ws["A2"] = "기준년월 : [2026.09]"
ws["A3"] = "단위 : 백만원 · 조회사무소 : 리스크관리부"
ws["A4"], ws["B4"], ws["C4"] = "상품코드", "상품명", "차주수"
ws.merge_cells("D4:F4")
ws["D4"] = "손실충당금"
ws["D5"], ws["E5"], ws["F5"] = "대출채권", "미수이자", "가지급금"
ws["G4"] = "충당부채"
ws["X4"] = "손실충당금 및 충당부채 총계"
for c in ("A4", "B4", "C4", "D4", "G4", "X4"):
    ws[c].font = Font(bold=True)
    ws[c].alignment = Alignment(horizontal="center", wrap_text=True)
r_no, delta = 6, Decimal(0)
total_prev = next((Decimal(str(r["provision"])) for r in prov if r.get("is_total")), None)
for r in sorted((x for x in prov if not x.get("is_total")), key=lambda x: x["product_code"]):
    old = Decimal(str(r["provision"]))
    g = rng.uniform(0.02, 0.07) if r["product_code"] in inc else rng.uniform(-0.01, 0.02)
    new = q(old * Decimal(str(1 + g)))
    delta += new - old
    loan, accr, adv = q(new * Decimal("0.92")), q(new * Decimal("0.05")), q(new * Decimal("0.01"))
    ws.cell(row=r_no, column=1, value=str(r["product_code"]))
    ws.cell(row=r_no, column=2, value=r.get("product_name"))
    ws.cell(row=r_no, column=3, value=int(old / 8) + rng.randint(-3, 12))
    ws.cell(row=r_no, column=4, value=float(loan))
    ws.cell(row=r_no, column=5, value=float(accr))
    ws.cell(row=r_no, column=6, value=float(adv))
    ws.cell(row=r_no, column=7, value=float(new - loan - accr - adv))
    ws.cell(row=r_no, column=24, value=float(new))
    r_no += 1
r_no += 1
ws.cell(row=r_no, column=1, value="합계")
ws.cell(row=r_no, column=24, value=float(q(total_prev * Decimal("1.004") + delta)) if total_prev is not None else None)
made.append((save(wb, "(202609)7790_상품별_충당금.xlsx"), "PROVISION_BY_PRODUCT", "bank"))

# ── 2. 신용평가모형 등급별 모니터링 집계 2026.09 (데이터셋 컬럼명 그대로)
mm = [r for r in rows("MODEL_MONITORING_GRADE") if str(r["as_of"]) == PREV_ASOF]
new_rows = []
for r in sorted(mm, key=lambda x: (x["model_code"], x["grade_seq"])):
    n = int(r["obs_count"] * (1 + rng.uniform(0.003, 0.012)))
    if (r["model_code"], r["grade"]) == ("CORP_MICRO", "6B"):
        n = 2490  # 시나리오: 6B 관측 2,472 → 2,490 · 부도 196 → 205 (실제부도율 8.233%)
    d = 205 if (r["model_code"], r["grade"]) == ("CORP_MICRO", "6B") else int(r["default_count"]) + rng.randint(0, max(1, int(r["default_count"] * 0.04)))
    new_rows.append({**r, "as_of": ASOF, "obs_count": n, "default_count": min(d, n), "exposure_amt": None if r.get("exposure_amt") is None else float(q(Decimal(str(r["exposure_amt"])) * Decimal(str(1 + rng.uniform(-0.01, 0.02)))))})
for model in {r["model_code"] for r in new_rows}:
    tot = sum(r["obs_count"] for r in new_rows if r["model_code"] == model)
    for r in new_rows:
        if r["model_code"] == model:
            r["obs_share"] = float(q(Decimal(r["obs_count"]) * 100 / Decimal(tot), "0.01"))
cols = ["as_of", "model_code", "grade", "grade_seq", "obs_count", "default_count", "obs_share", "ks", "sdr_rank", "exposure_amt"]
made.append((table("신용평가모형_등급별_모니터링_집계(2026.09).xlsx", "등급별", cols, [[r.get(c) for c in cols] for r in new_rows]), "MODEL_MONITORING_GRADE", "bank"))

# ── 3. Oblg_grd_def 202609 — 원본 컬럼명(대소문자 보존)
ob = [r for r in rows("OBLG_GRD_DEF") if str(r["std_ym"]) == PREV_YM]
labels = ["1", "2A", "2B", "3A", "3B", "4A", "4B", "5A", "5B", "5C", "6A", "6B", "6C", "7A", "7B", "8"]
out = []
for r in ob:
    g = int(r.get("Ls_nics_GRDC") or 0)
    if 1 <= g <= 16 and rng.random() < 0.015:
        g = max(1, min(16, g + rng.choice((-1, 1))))
    bal = Decimal(str(r["ln_bac"])) * Decimal(str(1 + rng.gauss(0, 0.01)))
    out.append({**r, "std_ym": YM, "Ls_nics_GRDC": g if r.get("Ls_nics_GRDC") is not None else None, "ln_bac": won10k(max(Decimal(0), bal))})
c = s.query(CalcRun).filter(CalcRun.metric_code == "br_prospect_pick").order_by(CalcRun.created_at.desc()).first()
picked = ((c.result or {}).get("selected") or [])[:2] if c else []
seq = max(int(str(r["brwr_no"])[-6:]) for r in ob if str(r["brwr_no"]).startswith("B1068")) + 1
onboard = []
for p in picked:
    onboard.append({"std_ym": YM, "brwr_no": f"B1068{seq:06d}", "brwr_nm": p["corp_nm"], "SUVMN_BRC": OFFICE, "Ls_nics_GRDC": labels.index(p["exp_grade"]) + 1,
                    "ln_bac": float(Decimal(str(p["loan_mn"])) * 1_000_000), "onlycard": 0, "STD_INDS_CFC": p["ind_cd"], "crdevl_pttp_dsc_f": str(p["nics_model"]), "default": 0})
    seq += 1
out += onboard
cols = ["std_ym", "brwr_no", "brwr_nm", "SUVMN_BRC", "Ls_nics_GRDC", "ln_bac", "onlycard", "STD_INDS_CFC", "crdevl_pttp_dsc_f", "default"]
made.append((table("Oblg_grd_def_202609.xlsx", "Oblg_grd_def", cols, [[r.get(k) for k in cols] for r in out]), "OBLG_GRD_DEF", "bank"))

# ── 4. TB_JC_PM_C_LA 202609 — 소매 익스포저(한글 별칭 헤더)
pools = {r["pool_id"]: r for r in rows("RETAIL_POOL_MASTER")}
re_ = [r for r in rows("RETAIL_EXPOSURE") if str(r["std_ym"]) == PREV_YM]
out = []
for r in re_:
    amt = Decimal(str(r["exp_amt"])) * Decimal(str(1 - rng.uniform(0, 0.012)))
    dlq = int(r.get("dlq_yn") or 0)
    if rng.random() < 0.004:
        dlq = 1 - dlq
    out.append({**r, "std_ym": YM, "exp_amt": won10k(amt), "dlq_yn": dlq})
seomin_pools = [pid for pid, p in pools.items() if p.get("seomin_eval")]
seq = max(int(str(r["brwr_id"])[-6:]) for r in re_ if str(r["brwr_id"]).startswith("R1068")) + 1
new_seomin = 8
for i in range(new_seomin):
    out.append({"std_ym": YM, "brwr_id": f"R1068{seq + i:06d}", "office_cd": OFFICE, "pool_id": seomin_pools[i % len(seomin_pools)], "exp_amt": float(rng.randint(25, 35) * 1_000_000), "cb_grade": rng.choice((6, 7, 7, 8)), "dlq_yn": 0})
hdr = ["기준년월", "차주식별번호", "사무소코드", "해당상품Pool", "여신익스포져금액", "CB등급", "연체여부"]
keys = ["std_ym", "brwr_id", "office_cd", "pool_id", "exp_amt", "cb_grade", "dlq_yn"]
made.append((table("TB_JC_PM_C_LA_202609.xlsx", "TB_JC_PM_C_LA", hdr, [[r.get(k) for k in keys] for r in out]), "RETAIL_EXPOSURE", "bank"))

# ── 5. 영업점 업적평가 득점 현황 2026.09.30 — 실적 반영(득점 = 실적 ÷ 목표 × 배점, 상한 배점)
def mv(code: str) -> Decimal:
    v = s.query(MetricValue).filter(MetricValue.metric_code == code).order_by(MetricValue.created_at.desc()).first()
    return Decimal(str(v.value))


sc = [r for r in rows("BRANCH_KPI_SCORE") if str(r["as_of"]) == PREV_ASOF and r["office_cd"] == OFFICE]
by = {r["item_cd"]: dict(r, as_of=ASOF) for r in sc}
prime_add = sum(Decimal(str(p["weighted_mn"])) for p in picked)
corp_add = sum(Decimal(str(p["loan_mn"])) for p in picked)
it = by["CORE_PRIME"]
it["actual_value"] = float(Decimal(str(it["actual_value"])) + prime_add)
it["score"] = float(q(min(Decimal(1), Decimal(str(it["actual_value"])) / mv("br_prime_target")) * mv("br_prime_points"), "0.01"))
it = by["CORE_CORP_LOAN"]
pts = q(Decimal(str(it["score"])) / (Decimal(str(it["actual_value"])) / Decimal(str(it["target_value"]))))  # 배점 역산(60점)
it["actual_value"] = float(Decimal(str(it["actual_value"])) + corp_add)
it["score"] = float(q(min(Decimal(1), Decimal(str(it["actual_value"])) / Decimal(str(it["target_value"]))) * pts, "0.01"))
it = by["INCL_SEOMIN"]
it["actual_value"] = int(it["actual_value"]) + new_seomin
it["score"] = float(q(min(Decimal(1), Decimal(it["actual_value"]) / mv("br_seomin_target")) * mv("br_seomin_points"), "0.01"))
for k in ("CUST_BASE", "CUST_DIGITAL"):
    by[k]["score"] = float(Decimal(str(by[k]["score"])) + rng.choice((0, 1, 1)))
by["BIZ_CORE"]["score"] = float(sum(Decimal(str(by[k]["score"])) for k in ("CORE_CORP_LOAN", "CORE_PRIME", "CORE_DEPOSIT")))
by["BIZ_INCL"]["score"] = float(sum(Decimal(str(by[k]["score"])) for k in ("INCL_PROD", "INCL_INCL", "INCL_SEOMIN")))
cols = ["office_cd", "period", "as_of", "level", "group_cd", "item_cd", "score", "actual_value", "target_value", "unit", "note"]
made.append((table("영업점 업적평가 득점 현황(202609).xlsx", "득점", cols, [[r.get(k) for k in cols] for r in by.values()]), "BRANCH_KPI_SCORE", "bank"))

# ── 6. PDF 2종 (나눔고딕 내장)
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

pdfmetrics.registerFont(TTFont("Nanum", "/usr/share/fonts/truetype/nanum/NanumGothic.ttf"))
H = ParagraphStyle("h", fontName="Nanum", fontSize=15, leading=22, spaceAfter=10)
H2 = ParagraphStyle("h2", fontName="Nanum", fontSize=12, leading=18, spaceBefore=8, spaceAfter=4)
P = ParagraphStyle("p", fontName="Nanum", fontSize=10.5, leading=17)


def pdf(name: str, title: str, pages: list[list[tuple[str, str]]]) -> str:
    path = os.path.join(OUT, name)
    story = [Paragraph(title, H), Paragraph("※ 시뮬레이션 시험 자료 — 실제 기관 문서가 아닙니다(플랫폼 업로드 시험용).", P), Spacer(1, 8)]
    for i, page in enumerate(pages):
        for kind, text in page:
            story.append(Paragraph(text, H2 if kind == "h" else P))
        if i < len(pages) - 1:
            story.append(PageBreak())
    SimpleDocTemplate(path, pagesize=A4, title=title, author="NH 플랫폼 시험").build(story)
    return path


made.append((pdf("(시뮬레이션) 2026년 하반기 서민금융 공급 확대 및 취약차주 보호 방안.pdf", "(시뮬레이션) 2026년 하반기 서민금융 공급 확대 및 취약차주 보호 방안", [
    [("h", "1. 추진 배경"), ("p", "고금리 장기화로 중·저신용 서민의 자금 수요가 늘었으나 제도권 금융 공급은 정체되어 있다. 정책서민금융 상품의 공급 목표를 상향하고, "
      "금융회사의 자체 서민금융 상품 취급을 확대해 불법사금융 이용을 줄인다."),
     ("h", "2. 주요 내용"), ("p", "(1) 새희망홀씨 등 은행권 자체 서민금융 공급 목표를 전년 대비 10% 이상 확대하고, 영업점 평가에 신규 공급 실적을 반영하도록 권고한다. "
      "(2) 동일 차주 연간 대출 한도는 3,500만원 이내로 운영한다. (3) 취약차주 연체 발생 시 채무조정 안내를 의무화하고, 연체 관리 실적을 별도 점검한다.")],
    [("h", "3. 금융회사 협조 사항"), ("p", "서민금융 공급 확대가 건전성에 미치는 영향을 정기적으로 점검하고, 상품별 충당금 비중과 연체율 추이를 반기마다 보고한다. "
      "포용금융 상품 충당금 비중이 크게 늘면 해당 Pool 의 연체 발생과 연체율을 집중 점검한다."),
     ("h", "4. 시행"), ("p", "2026년 10월 1일부터 시행하며, 하반기 실적부터 적용한다.")],
]), "doc:gov_policy", "ext"))
made.append((pdf("(시뮬레이션) 신용리스크관리부 공문 - 소기업 모형 6B 등급 부도율 점검 요청(2026.09).pdf", "(시뮬레이션) 소기업 모형 6B 등급 부도율 점검 및 등급계량화 재검토 요청", [
    [("h", "1. 개요"), ("p", "2026년 9월 말 기준 소기업 모형 6B 등급의 실제부도율이 예측부도율 대표값을 2개월 연속 상회했다. 최대값 이내이나 초과 폭이 커지는 추세여서 "
      "등급계량화(PD 매핑) 재검토 필요성을 점검한다."),
     ("h", "2. 요청 사항"), ("p", "(1) 6B 등급 신규 취급 시 재무 · 비재무 항목을 추가 확인한다. (2) 모형 모니터링 보고서에 등급별 실제부도율과 대표 · 최대 PD 비교표를 첨부한다. "
      "(3) PSI 는 관찰 구간 기준으로 계속 관리하고, 불안정 구간 진입 시 즉시 보고한다.")],
    [("h", "3. 일정"), ("p", "10월 중 등급계량화 재검토안을 마련하고, 11월 모형위원회에 부의한다.")],
]), "doc:model_method", "bank"))

print("\n".join(f"{os.path.basename(p)}\t{t}\t{o}\t{os.path.getsize(p):,}B" for p, t, o in made))
print(f"prospects onboarded: {[p['corp_nm'] for p in picked]} · seomin new {new_seomin}")
